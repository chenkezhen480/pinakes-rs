"""Writes the PDF test documents (tests/convert.rs, src/convert/pdf.rs): policy-zh.docx with
python-docx (headings, lists, ruled tables, a header and a footer with page numbers, in
Chinese), then policy-zh.pdf exported from it by Microsoft Word (Windows, through COM), and
scanned-zh.pdf, its first page as a picture only (no text layer), with PyMuPDF. The files are
committed; run this only to change them: `python tests/fixtures/make_pdf_fixtures.py`.
"""

import os
import pathlib
import subprocess

import pymupdf
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = pathlib.Path(os.path.dirname(os.path.abspath(__file__)))


class Doc:
    def __init__(self, zh=False):
        self.d = Document()
        if zh:
            for name in ["Normal", "Heading 1", "Heading 2", "Heading 3", "Title", "List Bullet", "List Number"]:
                st = self.d.styles[name]
                rpr = st.element.get_or_add_rPr()
                fonts = rpr.find(qn("w:rFonts"))
                if fonts is None:
                    fonts = OxmlElement("w:rFonts")
                    rpr.append(fonts)
                fonts.set(qn("w:eastAsia"), "SimSun")

    def heading(self, text, level):
        self.d.add_heading(text, level)

    def para(self, text):
        self.d.add_paragraph(text)

    def bullets(self, texts, numbered=False):
        for t in texts:
            self.d.add_paragraph(t, style="List Number" if numbered else "List Bullet")

    def table(self, rows, ruled=True):
        t = self.d.add_table(rows=len(rows), cols=len(rows[0]))
        if ruled:
            t.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, text in enumerate(row):
                t.cell(r, c).text = text

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def header_footer(self, header, footer):
        sec = self.d.sections[0]
        sec.header.paragraphs[0].text = header
        p = sec.footer.paragraphs[0]
        p.text = footer + " "
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "1"
        r.append(t)
        fld.append(r)
        p._p.append(fld)

    def save(self, name):
        self.d.save(HERE / f"{name}.docx")



def policy_zh():
    d = Doc(zh=True)
    d.header_footer("华东仓储有限公司 内部资料", "仓库作业规范 第")
    d.heading("仓库作业规范（2026 年版）", 0)
    d.para("本规范适用于华东区所有自营仓库，自 2026 年 4 月 1 日起执行，原 2024 年版同时废止。")
    d.heading("一、入库", 1)
    d.para("到货车辆须提前 24 小时预约，未预约车辆排在当日最后卸货。收货员按采购单逐箱扫码，差异超过 2% 时暂停收货并通知采购。")
    d.bullets([
        "冷链货物到货温度须在 2℃ 至 8℃ 之间，超出范围一律拒收。",
        "外包装破损的箱子单独拍照，并在系统中标记为待检。",
        "危险品（如锂电池 SKU-BAT-07）须放入隔离区，由持证人员处理。",
    ])
    d.heading("二、库存与盘点", 1)
    d.para("每月最后一个周日进行全盘，每周三对 A 类货物抽盘。盘点差异的处理标准如下表。")
    d.table([
        ["货物类别", "允许差异", "处理时限", "审批人"],
        ["A 类（高价值）", "0", "当日", "仓库经理"],
        ["B 类", "0.5%", "3 个工作日", "主管"],
        ["C 类", "1%", "7 个工作日", "组长"],
    ])
    d.heading("2.1 呆滞库存", 2)
    d.para("连续 90 天无出库的货物视为呆滞库存，由计划部每月 5 日前出具处理方案；单个 SKU 金额超过 5 万元的，须报区域总监审批。")
    d.page_break()
    d.heading("三、出库与配送", 1)
    d.para("订单在 15:00 前确认的当日发出，之后确认的次日发出。大件商品（单件超过 30 公斤）由合作承运商上门提货。")
    d.bullets([
        "拣货完成后须二次复核，复核人与拣货人不得为同一人。",
        "出库单打印后 30 分钟内未装车的，须重新扫码确认。",
        "跨境订单须附中英文装箱单（Packing List）。",
    ], numbered=True)
    d.table([
        ["重量区间", "承运商", "费用（元/单）"],
        ["30 公斤以下", "顺达快运", "免费"],
        ["30 至 70 公斤", "重货通", "45"],
        ["70 公斤以上", "重货通（双人）", "95"],
    ])
    d.heading("四、安全", 1)
    d.para("叉车须持证驾驶，库内限速每小时 5 公里。发生事故后 1 小时内上报安全员，并在系统中填写事故单。")
    d.save("policy-zh")


def export_with_word(docx, pdf):
    """Word's own PDF export (`ExportAsFixedFormat`, format 17), as people's PDFs come."""
    script = (
        "$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
        f"$d = $w.Documents.Open('{docx}', $false, $true); $d.ExportAsFixedFormat('{pdf}', 17); "
        "$d.Close($false); try { $w.Quit() } catch {}"
    )
    subprocess.run(["powershell", "-NoProfile", "-Command", script], check=True)


def scan_of_first_page(pdf, out):
    src = pymupdf.open(pdf)
    pix = src[0].get_pixmap(dpi=100, colorspace=pymupdf.csGRAY)
    doc = pymupdf.open()
    page = doc.new_page(width=src[0].rect.width, height=src[0].rect.height)
    page.insert_image(page.rect, stream=pix.tobytes("png"))
    doc.save(out, deflate=True)


policy_zh()
export_with_word(HERE / "policy-zh.docx", HERE / "policy-zh.pdf")
scan_of_first_page(HERE / "policy-zh.pdf", HERE / "scanned-zh.pdf")
