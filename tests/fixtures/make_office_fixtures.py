"""Writes the Office test documents the converter tests read (tests/convert.rs): Word, Excel
and PowerPoint files in Chinese and English, made with python-docx, openpyxl and python-pptx.
The files are committed; run this only to change them: `python tests/fixtures/make_office_fixtures.py`.
"""

import datetime
import os

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

HERE = os.path.dirname(os.path.abspath(__file__))


def link(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def returns_policy():
    d = Document()
    d.add_heading("Returns policy", 0)
    d.add_heading("Who may return goods", 1)
    p = d.add_paragraph("Members may return goods within ")
    p.add_run("seven days").bold = True
    p.add_run(" of delivery; see ")
    link(p, "the full terms", "https://example.com/terms")
    p.add_run(".")
    d.add_paragraph("Unopened goods", style="List Bullet")
    d.add_paragraph("Goods with their original receipt", style="List Bullet")
    d.add_paragraph("Bought online", style="List Bullet 2")
    d.add_heading("How to return", 2)
    d.add_paragraph("Ask for a return label in the app.", style="List Number")
    d.add_paragraph("Pack the goods with the receipt.", style="List Number")
    d.add_paragraph("Hand the parcel to the courier.", style="List Number")
    d.add_heading("Refund times", 1)
    t = d.add_table(rows=1, cols=3)
    t.rows[0].cells[0].text = "Payment"
    t.rows[0].cells[1].text = "Refund within"
    t.rows[0].cells[2].text = "Note"
    for row in [
        ("Card", "5 days", "to the same card"),
        ("Wallet", "1 day", "instant for gold members"),
        ("Cash on delivery", "7 days", "by bank transfer"),
    ]:
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            c.text = v
    d.add_paragraph("Fresh food cannot be returned.")
    d.save(os.path.join(HERE, "returns-policy.docx"))


def warehouse_rules():
    d = Document()
    d.add_heading("仓储作业规范", 0)
    d.add_heading("一、入库", 1)
    d.add_paragraph("到货后两小时内完成验收，验收不合格的货品放入隔离区。")
    d.add_paragraph("核对送货单与采购单", style="List Number")
    d.add_paragraph("检查外包装与数量", style="List Number")
    d.add_paragraph("录入仓储系统", style="List Number")
    d.add_heading("二、盘点", 1)
    d.add_heading("2.1 周期", 2)
    p = d.add_paragraph("每月最后一个周五进行")
    p.add_run("全盘").bold = True
    p.add_run("，差异超过千分之三须复盘。")
    d.add_heading("2.2 责任人", 2)
    t = d.add_table(rows=1, cols=3)
    for c, v in zip(t.rows[0].cells, ["区域", "责任人", "复核"]):
        c.text = v
    for row in [("A 区（食品）", "王强", "李娜"), ("B 区（日化）", "陈晨", "李娜"), ("冷库", "赵磊", "周敏")]:
        for c, v in zip(t.add_row().cells, row):
            c.text = v
    d.add_paragraph("温控", style="List Bullet")
    d.add_paragraph("冷库保持零下十八度以下", style="List Bullet 2")
    d.add_paragraph("每四小时记录一次", style="List Bullet 2")
    d.save(os.path.join(HERE, "warehouse-rules-zh.docx"))


def inventory():
    wb = Workbook()
    ws = wb.active
    ws.title = "库存"
    ws.append(["品名", "", "仓库", "数量", "入库日期"])
    ws.append(["矿泉水 550ml", "", "A 区", 1200, datetime.date(2026, 9, 1)])
    ws.append([])
    ws.append(["洗衣液 2L", "", "B 区", 340, datetime.date(2026, 9, 15)])
    ws.append(["冷冻水饺", "", "冷库", 85.5, datetime.date(2026, 10, 2)])
    ws2 = wb.create_sheet("Suppliers")
    ws2.append(["Supplier", "Contact", "Lead time (days)"])
    ws2.append(["Blue Spring Ltd", "ops@bluespring.example", 3])
    ws2.append(["CleanCo", "sales@cleanco.example", 7])
    wb.create_sheet("Empty")
    wb.save(os.path.join(HERE, "inventory.xlsx"))


def deck(path, title, subtitle, bullets, table, notes):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = title
    s.placeholders[1].text = subtitle
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = bullets[0]
    body = s.placeholders[1].text_frame
    body.text = bullets[1][0][0]
    for text, level in bullets[1][1:]:
        p = body.add_paragraph()
        p.text = text
        p.level = level
    s.notes_slide.notes_text_frame.text = notes
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = table[0]
    rows = table[1]
    shape = s.shapes.add_table(len(rows), len(rows[0]), Inches(1), Inches(2), Inches(8), Inches(2))
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            shape.table.cell(i, j).text = v
    prs.save(os.path.join(HERE, path))


def decks():
    deck(
        "quarterly-review-zh.pptx",
        "第三季度经营回顾",
        "华东大区 · 2026 年 10 月",
        (
            "销售亮点",
            [("销售额同比增长 18%", 0), ("华东贡献过半", 1), ("新客占比六成", 1), ("退货率降至 2.1%", 0)],
        ),
        ("各区销量", [["区域", "销量（万件）", "同比"], ["华东", "52", "+21%"], ["华南", "31", "+9%"]]),
        "重点讲华东的增长原因：会员日活动。",
    )
    deck(
        "product-launch.pptx",
        "Launching the Gold membership",
        "Product team, October 2026",
        (
            "What members get",
            [("Two free shipping coupons a month", 0), ("Used on orders over 49 yuan", 1), ("Priority support", 0)],
        ),
        ("Prices", [["Plan", "Price", "Coupons"], ["Silver", "0", "0"], ["Gold", "99 a year", "24"]]),
        "Stress that coupons do not roll over.",
    )


if __name__ == "__main__":
    returns_policy()
    warehouse_rules()
    inventory()
    decks()
